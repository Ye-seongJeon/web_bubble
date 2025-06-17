import os
from openpyxl import Workbook
from OCC.Core.BRepGProp import brepgprop_VolumeProperties
from OCC.Core.GProp import GProp_GProps
from OCC.Extend.DataExchange import read_step_file
from OCC.Core.Bnd import Bnd_Box
from OCC.Core.BRepBndLib import brepbndlib_Add


import math

# stl_file_list = os.listdir(r"models\student\stl")
# stp_file_list = os.listdir(r"models\student\stp")
# CATPart_file_list = os.listdir(r"models\student\CATPart")

# print("stl 검사")
# for i in range(len(stl_file_list)):
#     stl_fileExtension = os.path.splitext(stl_file_list[i])[1]
#     if stl_fileExtension == '.stl':
#         print("True")
#     else:
#         print("{} 의 stl 파일의 확장자가 잘못되었습니다.".format(stl_file_list[i]))

# print("stp 검사")
# for i in range(len(stp_file_list)):
#     stp_fileExtension = os.path.splitext(stp_file_list[i])[1]
#     if stp_fileExtension == '.stp':
#         print("True")
#     else:
#         print("{} 의 stp 파일의 확장자가 잘못되었습니다.".format(stp_file_list[i]))
class FileExtension:
    def __init__(self, format_name):
        self.format = format_name
    def CheckList(self, format_name, file_list):
        bool_list =[]
        for i in range(len(file_list)):
            fileExtension = os.path.splitext(file_list[i])[1]
            if fileExtension == "."+format_name:
                bool_list.append(True)
            else:
                # bool_list.append("{0} 파일의 확장자가 잘못되었습니다.".format(file_list[i]))
                bool_list.append(False)
        return bool_list
    # def score(self, bool_list):
    #     count = 0
    #     for value in bool_list:
    #         if value == True:
    #             count += 1
    #     score = count/len(bool_list)
    #     return score

class VolumeCheck:
    def __init__(self,stp_path,stp_file_list):
        shps = []
        stp = stp_file_list
        stp_student = read_step_file(stp_path+r"\\"+stp)
        shps.append(stp_student)

        self.shapes = shps

    def measure_shape_volume(self,shape):
        """Returns shape volume"""
        inertia_props = GProp_GProps()
        brepgprop_VolumeProperties(shape, inertia_props)
        mass = inertia_props.Mass()
        return mass
    
    def measure_bnd_box(self, shape):
        bbox = Bnd_Box()
        brepbndlib_Add(shape, bbox) # shape에 bbox 추가
        xmin, ymin, zmin, xmax, ymax, zmax = bbox.Get()
        x_length = xmax - xmin
        x_length = round(x_length, 3)
        y_length = ymax - ymin
        y_length = round(y_length, 3)
        z_length = zmax - zmin
        z_length = round(z_length, 3)
        bnd_list = [x_length, y_length, z_length]
        bnd_list.sort()
        return bnd_list

    def VolScores(self,T_stp_path, T_stp_file):
        stp_T = read_step_file(T_stp_path+r"\\"+T_stp_file)
        vol_Tshp = self.measure_shape_volume(stp_T)
        for shp in self.shapes:
            # 1 볼륨 계산
            vol_shp = self.measure_shape_volume(shp)

            # 2배 초과 = 0점
            if (vol_shp/vol_Tshp > 2 or vol_shp == 0):
                VolScores = 0
            # 0~1배 (작을 경우)
            elif (vol_shp/vol_Tshp < 1):
                VolScores = 100*(vol_shp/vol_Tshp)
            # 1~2배 (클 경우)
            else:
                VolScores = -100*(vol_shp/vol_Tshp) + 200
        print("vol_shp",vol_shp)
        print("vol_Tshp",vol_Tshp)
        return VolScores
    
    def BndScores(self,T_stp_path, T_stp_file):
        stp_T = read_step_file(T_stp_path+r"\\"+T_stp_file)
        teacher_bnd_list = self.measure_bnd_box(stp_T)
        for shp in self.shapes:
            # 2 바운딩박스 치수 계산
            BndScores = 0
            student_bnd_list = self.measure_bnd_box(shp)
            for i in range(3):
        #         if (teacher_bnd_list[i] == student_bnd_list[i]):
        #             BndScores += 33.33
        

                if (student_bnd_list[i]/teacher_bnd_list[i] > 2):
                    BndScores += 0
                # 0~1배 (작을 경우)
                elif (student_bnd_list[i]/teacher_bnd_list[i] <= 1):
                    BndScores += 33.3*(student_bnd_list[i]/teacher_bnd_list[i])
                # 1~2배 (클 경우)
                else:
                    BndScores += -33.3*(student_bnd_list[i]/teacher_bnd_list[i]) + 66.6
            # Scores = round((VolScores + BndScores) / 2)
        BndScores = round(BndScores, 2)
        if BndScores == 99.9:
            BndScores = 100
        return BndScores





if __name__ =="__main__":
    VolScores = []
    vol_Tshp = 1
    shps = [0.05,0.1,0.2,0.5,1,2,5,10,20] # [34.94850021680094, 50.0, 65.05149978319906, 84.94850021680094, 100.0, 84.94850021680094, 65.05149978319906, 50.0, 34.94850021680094]
    for shp in shps:
        vol_shp = shp
        #log_vol_stu = math.log(vol_shp)
        score = 50 * (2-abs(math.log10(vol_shp/vol_Tshp))) # 0 ~ 100점 로그 스케일 비교 
        VolScores.append(score)
    print(VolScores)